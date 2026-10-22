"""Preview slides of a pptx without LibreOffice: freeform shapes (custGeom), text boxes and pictures, drawn in
PIL from the slide XML. Good enough to check geometry, stacking and text placement of decks made here.

    .venv-analysis/bin/python figures/deck/preview_pptx.py deck.pptx out_prefix [dpi]
"""
import io
import sys

import numpy as np
from lxml import etree
from matplotlib import font_manager as fm
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EMU = 914400
FONTS = {(False, False): fm.findfont("DejaVu Sans"),
         (True, False): fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="bold")),
         (False, True): fm.findfont(fm.FontProperties(family="DejaVu Sans", style="italic")),
         (True, True): fm.findfont(fm.FontProperties(family="DejaVu Sans", weight="bold", style="italic"))}


def color(el):
    c = el.find(f"{A}srgbClr") if el is not None else None
    if c is None:
        return None
    v = c.get("val")
    a = c.find(f"{A}alpha")
    al = int(a.get("val")) / 100000 if a is not None else 1.0
    return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4)) + (int(al * 255),)


def bez(p0, p1, p2, p3, n=12):
    t = np.linspace(0, 1, n)[:, None]
    return ((1 - t) ** 3 * p0 + 3 * (1 - t) ** 2 * t * p1 + 3 * (1 - t) * t ** 2 * p2 + t ** 3 * p3)


def render(deck, prefix, dpi=110):
    prs = Presentation(deck)
    W, H = int(prs.slide_width / EMU * dpi), int(prs.slide_height / EMU * dpi)
    s2 = dpi / EMU
    for i, slide in enumerate(prs.slides, 1):
        bg = slide.background.fill
        try:
            base = tuple(bg.fore_color.rgb) + (255,)
        except Exception:
            base = (255, 255, 255, 255)
        im = Image.new("RGBA", (W, H), base)
        for sh in slide.shapes._spTree.iterchildren():
            tag = etree.QName(sh).localname
            if tag == "pic":
                off = sh.find(f".//{A}off")
                ext = sh.find(f".//{A}ext")
                blip = sh.find(f".//{A}blip")
                rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                data = slide.part.related_part(rid).blob
                pic = Image.open(io.BytesIO(data)).convert("RGBA")
                x, y = int(int(off.get("x")) * s2), int(int(off.get("y")) * s2)
                w, h = max(1, int(int(ext.get("cx")) * s2)), max(1, int(int(ext.get("cy")) * s2))
                pic = pic.resize((w, h))
                im.alpha_composite(pic, (max(0, x), max(0, y)))
                continue
            if tag != "sp":
                continue
            off = sh.find(f".//{A}xfrm/{A}off")
            ext = sh.find(f".//{A}xfrm/{A}ext")
            if off is None:
                continue
            ox, oy = int(off.get("x")), int(off.get("y"))
            cx, cy = int(ext.get("cx")), int(ext.get("cy"))
            spPr = sh.find(f"{P}spPr")
            geom = spPr.find(f"{A}custGeom")
            layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
            d = ImageDraw.Draw(layer)
            if geom is not None:
                fill = spPr.find(f"{A}solidFill")
                patt = spPr.find(f"{A}pattFill")
                fc = color(fill) if fill is not None else (color(patt.find(f"{A}fgClr")) if patt is not None else None)
                ln = spPr.find(f"{A}ln")
                lc = color(ln.find(f"{A}solidFill")) if ln is not None and ln.find(f"{A}solidFill") is not None else None
                lw = max(1, int(int(ln.get("w", 12700)) * s2)) if lc else 0
                for path in geom.iter(f"{A}path"):
                    pw, ph = int(path.get("w")), int(path.get("h"))
                    sx, sy = cx / max(pw, 1), cy / max(ph, 1)
                    polys, cur = [], []
                    last = None
                    for c in path:
                        n = etree.QName(c).localname
                        pts = [np.array([ox + int(p.get("x")) * sx, oy + int(p.get("y")) * sy]) * s2
                               for p in c.iter(f"{A}pt")]
                        if n == "moveTo":
                            if len(cur) > 1:
                                polys.append((cur, False))
                            cur = [pts[0]]
                        elif n == "lnTo":
                            cur += pts
                        elif n == "cubicBezTo" and cur:
                            cur += list(bez(cur[-1], *pts))
                        elif n == "quadBezTo" and cur:
                            p0, p1, p2 = cur[-1], *pts
                            cur += list(bez(p0, p0 + 2 / 3 * (p1 - p0), p2 + 2 / 3 * (p1 - p2), p2))
                        elif n == "close":
                            polys.append((cur, True))
                            cur = []
                    if len(cur) > 1:
                        polys.append((cur, False))
                    for poly, closed in polys:
                        xy = [tuple(p) for p in poly]
                        if fc and path.get("fill") != "none" and len(xy) > 2:
                            d.polygon(xy, fill=fc)
                        if lc:
                            d.line(xy + ([xy[0]] if closed else []), fill=lc, width=lw)
            tx = sh.find(f"{P}txBody")
            if tx is not None:
                y = oy * s2
                for p in tx.iter(f"{A}p"):
                    algn = (p.find(f"{A}pPr").get("algn") if p.find(f"{A}pPr") is not None else "l") or "l"
                    runs = []
                    for r in p.iter(f"{A}r"):
                        rp = r.find(f"{A}rPr")
                        sz = int(rp.get("sz", "1800")) / 100 * dpi / 72
                        base = int(rp.get("baseline", "0"))
                        f = ImageFont.truetype(FONTS[(rp.get("b") == "1", rp.get("i") == "1")],
                                               max(4, int(sz * (0.7 if base else 1))))
                        runs.append((r.find(f"{A}t").text or "", f, color(rp.find(f"{A}solidFill")) or (0, 0, 0, 255),
                                     sz, base))
                    if not runs:
                        continue
                    tw = sum(d.textlength(t, font=f) for t, f, *_ in runs)
                    size = max(r[3] for r in runs)
                    x = {"l": ox * s2, "ctr": (ox + cx / 2) * s2 - tw / 2, "r": (ox + cx) * s2 - tw}.get(algn, ox * s2)
                    for t, f, c, sz, base in runs:
                        d.text((x, y + size * 0.905 - base / 100000 * size), t, font=f, fill=c, anchor="ls")
                        x += d.textlength(t, font=f)
                    y += size * 1.2
            im.alpha_composite(layer)
        im.convert("RGB").save(f"{prefix}_{i}.png")
        print(f"{prefix}_{i}.png")


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else 110)
