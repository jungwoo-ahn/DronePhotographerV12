"""Editable deck of the paper figures: every box, line, arrow and label is a native PowerPoint object.

    .venv-analysis/bin/python figures/deck/make_editable_deck.py   ->  figures/deck/paper_figures_editable.pptx

Each figure script's build() is drawn through pptx_export.SlideRenderer instead of being rasterised, so the
slides can be edited in place (text retyped, colours changed, shapes moved). Photographs and renders stay
pictures. Fonts are mapped to Arial / Courier New / Times New Roman (math).
"""
import importlib.util
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

FIG = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(FIG))
sys.path.insert(0, str(FIG / "overview"))
from pptx_export import add_figure  # noqa: E402


def load(rel):
    spec = importlib.util.spec_from_file_location(Path(rel).stem, FIG / rel)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


SLIDES = [
    ("Fig. 2 — Method overview", "overview/make_fig2.py", [],
     "figures/overview/make_fig2.py. Rollout in (c): All-People-Are-Sisters p00 s31 g02, d 2.01 -> 0.066, shoot at "
     "step 7 of chunk 4. The user sentence in (c) is illustrative."),
    ("Data-generation pipeline", "datagen/make_fig_datagen.py", [],
     "figures/datagen/make_fig_datagen.py, panel (b) schematic. Counts taken from the data on 2026-10-14."),
    ("Data-generation pipeline — real floor grid", "datagen/make_fig_datagen.py", ["--b", "real"],
     "figures/datagen/make_fig_datagen.py --b real. Floor grid re-run without the Gemini interest region "
     "(sbatch_placement_grid.sh, job 192374); VLM view scores 3, 2, 1, 8, 4 are from the placement run."),
    ("Qualitative rollouts (1/2)", "qualitative/make_fig_rollouts.py", ["--rows", "0,1,2,4"],
     "figures/qualitative/make_fig_rollouts.py --rows 0,1,2,4. Selected examples: Ours succeeds, pi0.5 / GR00T / "
     "T2A / DP all miss; image-to-action left out (it is given the GT goal image)."),
    ("Qualitative rollouts (2/2)", "qualitative/make_fig_rollouts.py", ["--rows", "3,5,6,7"],
     "figures/qualitative/make_fig_rollouts.py --rows 3,5,6,7."),
]


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0x1F, 0x21, 0x26)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = "Paper figures — editable"
    p.runs[0].font.size, p.runs[0].font.bold = Pt(44), True
    p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p2 = tb.text_frame.add_paragraph()
    p2.text = "Every box, arrow and label is a PowerPoint object; photographs and renders are pictures"
    p2.runs[0].font.size = Pt(18)
    p2.runs[0].font.color.rgb = RGBColor(0xD5, 0xD8, 0xDD)

    for title, rel, argv, notes in SLIDES:
        mod = load(rel)
        cv = mod.build(mod.args(argv))
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.5))
        tb.text_frame.text = title
        r = tb.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(22), True, "Arial"
        r.font.color.rgb = RGBColor(0x1F, 0x21, 0x26)
        n = add_figure(s, cv.fig, 0.35, 0.78, 12.63, 6.55)
        s.notes_slide.notes_text_frame.text = notes
        plt.close(cv.fig)
        print(f"{title}: {n}")

    out = Path(__file__).resolve().parent / "paper_figures_editable.pptx"
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
